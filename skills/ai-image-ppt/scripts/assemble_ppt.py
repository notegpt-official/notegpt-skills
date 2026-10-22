#!/usr/bin/env python3
"""
PPT assembly script

Assemble slide images from a directory into a PowerPoint presentation.
Each image becomes one slide and fills the entire canvas.
"""

import argparse
import os
import re
import sys
import tempfile
from typing import Dict, List, Optional, Tuple

SUPPORTED_ASPECT_RATIOS = ("16:9", "4:3", "1:1", "3:4", "9:16", "21:9")


def dependency_hint() -> str:
    runtime_home = os.path.expanduser(os.environ.get("IMAGE_PPT_HOME", "~/.image-ppt-skill"))
    python = os.path.join(
        runtime_home,
        ".venv",
        "Scripts" if os.name == "nt" else "bin",
        "python.exe" if os.name == "nt" else "python",
    )
    runtime_script = os.path.join(
        os.path.dirname(os.path.abspath(__file__)),
        "image_ppt_runtime.py",
    )
    return (
        f"Run: python3 {runtime_script} bootstrap\n"
        f"Or run directly: {python} -m pip install -r "
        f"{os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'requirements.txt')}"
    )


def get_slide_images(ppt_project_dir: str) -> List[str]:
    """
    Get slide image paths sorted by slide number.

    Args:
        ppt_project_dir: PPT project directory containing the origin_image subdirectory

    Returns:
        Image file paths in slide order
    """
    # Read from the origin_image subdirectory
    origin_image_dir = os.path.join(ppt_project_dir, "origin_image")

    if not os.path.exists(origin_image_dir):
        print(f"Error: origin_image directory does not exist: {origin_image_dir}")
        return []

    print(f"Reading images from origin_image: {origin_image_dir}")

    # Supported image formats
    image_extensions = {'.png', '.jpg', '.jpeg', '.gif', '.bmp'}

    # Include only final slide images, excluding sample_slide.png, drafts, and references.
    slide_name_pattern = re.compile(r"^slide_(\d+)\.(png|jpe?g|gif|bmp)$", re.IGNORECASE)
    image_files = []
    for file in os.listdir(origin_image_dir):
        file_path = os.path.join(origin_image_dir, file)
        if os.path.isfile(file_path):
            ext = os.path.splitext(file)[1].lower()
            if ext in image_extensions and slide_name_pattern.match(file):
                image_files.append(file_path)

    def slide_sort_key(path: str) -> Tuple[int, str]:
        filename = os.path.basename(path)
        stem = os.path.splitext(filename)[0]
        return (int(re.search(r"^slide_(\d+)", stem).group(1)), filename.lower())

    # Sort numerically so slide_2 precedes slide_10.
    # Use filenames as a stable tie-breaker for equal slide numbers.
    image_files.sort(key=slide_sort_key)

    return image_files


def load_speaker_notes(ppt_project_dir: str) -> Dict[int, str]:
    """
    Load per-slide speaker notes from speech.md.

    Supported heading formats:
    - ## Slide 1: Title
    - ### Slide 1: Title
    - Chinese numbered slide headings are also supported.
    """
    speech_path = os.path.join(ppt_project_dir, "speech.md")
    if not os.path.exists(speech_path):
        return {}

    with open(speech_path, "r", encoding="utf-8") as f:
        content = f.read()

    notes: Dict[int, str] = {}
    current_slide: Optional[int] = None
    current_lines: List[str] = []
    heading_pattern = re.compile(r"^#{2,4}\s*(?:Slide\s*(\d+)|\u7b2c\s*(\d+)\s*\u9875)\b.*$", re.IGNORECASE)

    def flush_current() -> None:
        if current_slide is None:
            return
        text = "\n".join(current_lines).strip()
        if text:
            notes[current_slide] = text

    for line in content.splitlines():
        match = heading_pattern.match(line.strip())
        if match:
            flush_current()
            current_slide = int(match.group(1) or match.group(2))
            current_lines = []
            continue

        if current_slide is not None:
            current_lines.append(line)

    flush_current()
    return notes


def compress_image_if_needed(
    image_path: str,
    max_size_mb: float = 2.0,
    quality_step: int = 5
) -> Optional[str]:
    """
    Compress oversized images and return the temporary file path.

    Args:
        image_path: Original image path
        max_size_mb: Maximum file size in MB
        quality_step: Quality decrement per compression attempt

    Returns:
        str: Temporary file path if compression is needed; otherwise None
    """
    max_size_bytes = max_size_mb * 1024 * 1024

    # Check the original file size
    file_size = os.path.getsize(image_path)

    if file_size <= max_size_bytes:
        # No compression needed
        return None

    print(f"  Image size {file_size / 1024 / 1024:.2f}MB, compression needed...")

    try:
        from PIL import Image

        # Open the image
        img = Image.open(image_path)

        # Convert RGBA to RGB for JPEG output
        if img.mode in ('RGBA', 'LA', 'P'):
            background = Image.new('RGB', img.size, (255, 255, 255))
            if img.mode == 'P':
                img = img.convert('RGBA')
            background.paste(img, mask=img.split()[-1] if img.mode in ('RGBA', 'LA') else None)
            img = background

        # Create a temporary file
        temp_fd, temp_path = tempfile.mkstemp(suffix='.jpg')
        os.close(temp_fd)

        # Start compression at high quality
        quality = 95
        while quality > 20:
            img.save(temp_path, 'JPEG', quality=quality, optimize=True)
            compressed_size = os.path.getsize(temp_path)

            if compressed_size <= max_size_bytes:
                print(f"  Compressed successfully: {compressed_size / 1024 / 1024:.2f}MB (quality: {quality})")
                return temp_path

            quality -= quality_step

        # Try resizing if the image is still too large
        print(f"  Quality reduction was insufficient; trying smaller dimensions...")
        scale = 0.9
        while scale > 0.3:
            new_width = int(img.width * scale)
            new_height = int(img.height * scale)
            resized_img = img.resize((new_width, new_height), Image.Resampling.LANCZOS)

            resized_img.save(temp_path, 'JPEG', quality=85, optimize=True)
            compressed_size = os.path.getsize(temp_path)

            if compressed_size <= max_size_bytes:
                print(f"  Compressed successfully: {compressed_size / 1024 / 1024:.2f}MB (scale: {scale:.0%})")
                return temp_path

            scale -= 0.1

        # Return the last result if the size limit cannot be reached
        print(f"  Warning: could not compress below {max_size_mb}MB; using the smallest version")
        return temp_path

    except ImportError:
        print("Error: Pillow is not installed")
        print(dependency_hint())
        return None
    except Exception as e:
        print(f"  Warning: image compression failed: {e}")
        if 'temp_path' in locals() and os.path.exists(temp_path):
            os.remove(temp_path)
        return None


def create_presentation(
    image_files: List[str],
    output_path: str,
    aspect_ratio: str = "16:9",
    speaker_notes: Optional[Dict[int, str]] = None,
) -> bool:
    """
    Create a PowerPoint presentation.

    Args:
        image_files: List of slide image paths
        output_path: Output PPT file path
        aspect_ratio: Slide aspect ratio (default: 16:9; also 4:3, 1:1, 3:4, 9:16, 21:9)

    Returns:
        bool: True on success, False on failure
    """
    try:
        try:
            from pptx import Presentation
            from pptx.util import Inches
        except ImportError:
            print("Error: python-pptx is not installed")
            print(dependency_hint())
            return False

        # Create the presentation
        prs = Presentation()

        # Set slide dimensions
        if aspect_ratio not in SUPPORTED_ASPECT_RATIOS:
            print(f"Warning: unsupported aspect ratio {aspect_ratio}; using the default 16:9")
            aspect_ratio = "16:9"
        ratio_width, ratio_height = map(int, aspect_ratio.split(":"))
        prs.slide_width = Inches(10)
        prs.slide_height = Inches(10 * ratio_height / ratio_width)

        speaker_notes = speaker_notes or {}

        # Add each slide
        temp_files_to_cleanup = []
        for i, image_path in enumerate(image_files, 1):
            if not os.path.exists(image_path):
                print(f"Warning: image file does not exist: {image_path}")
                continue

            # Compress the image if needed
            compressed_path = compress_image_if_needed(image_path, max_size_mb=2.0)

            # Use the compressed image or the original
            image_to_use = compressed_path if compressed_path else image_path

            # Track temporary files for cleanup
            if compressed_path:
                temp_files_to_cleanup.append(compressed_path)

            # Use the blank layout (index 6)
            blank_slide_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_slide_layout)

            # Add the image to fill the entire slide
            slide.shapes.add_picture(
                image_to_use,
                left=0,
                top=0,
                width=prs.slide_width,
                height=prs.slide_height
            )

            note_text = speaker_notes.get(i)
            if note_text:
                notes_frame = slide.notes_slide.notes_text_frame
                notes_frame.clear()
                notes_frame.text = note_text

            print(f"✓ Added slide {i}: {os.path.basename(image_path)}")

        # Ensure the output directory exists
        output_dir = os.path.dirname(output_path)
        if output_dir and not os.path.exists(output_dir):
            os.makedirs(output_dir)

        # Save the presentation
        prs.save(output_path)
        print(f"\n✓ PPT saved: {output_path}")
        print(f"  Total slides: {len(image_files)}")
        if speaker_notes:
            matched_notes = sum(1 for i in range(1, len(image_files) + 1) if speaker_notes.get(i))
            print(f"  Notes written: {matched_notes}/{len(image_files)} slides")

        # Clean up temporary files
        for temp_file in temp_files_to_cleanup:
            try:
                if os.path.exists(temp_file):
                    os.remove(temp_file)
            except Exception as e:
                print(f"Warning: failed to clean up temporary file {temp_file}: {e}")

        return True

    except Exception as e:
        print(f"Error: failed to create PPT: {e}")
        import traceback
        traceback.print_exc()

        # Clean up temporary files even on failure
        if 'temp_files_to_cleanup' in locals():
            for temp_file in temp_files_to_cleanup:
                try:
                    if os.path.exists(temp_file):
                        os.remove(temp_file)
                except:
                    pass

        return False


def main():
    parser = argparse.ArgumentParser(
        description='Assemble slide images into a PowerPoint presentation',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='''
Examples:
  # Basic usage
  # Find the MyPresentation/ directory under /path/to/base/
  # Read images from MyPresentation/origin_image/
  # Save the PPT to MyPresentation/MyPresentation.pptx
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx

  # Use a 4:3 aspect ratio
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx --aspect-ratio 4:3

  # Initialize directories without generating a PPT
  python assemble_ppt.py /path/to/base/ MyPresentation.pptx --init

Required directory structure:
  /path/to/base/MyPresentation/
  ├── origin_image/
  │   ├── slide_01.png
  │   ├── slide_02.png
  │   └── ...
  ├── speech.md (optional; written to PPT speaker notes)
  └── MyPresentation.pptx (generated here)

Notes:
  - Place image files in the origin_image subdirectory
  - Read only final slide images such as slide_01.png and slide_02.png; ignore other images
  - Sort images by slide number
  - Recommended image filenames: slide_01.png, slide_02.png, ...
  - Each image fills the entire slide
  - If speech.md exists in the project directory, map Slide N headings to slide notes
        '''
    )

    parser.add_argument('base_dir', help='Base directory (parent of the PPT project directory)')
    parser.add_argument('output', help='Output PPT filename (.pptx)')
    parser.add_argument('--aspect-ratio', '--ar',
                        choices=SUPPORTED_ASPECT_RATIOS,
                        default='16:9',
                        help='Slide aspect ratio (default: 16:9)')
    parser.add_argument('--init',
                        action='store_true',
                        help='Create only the PPT project and origin_image directories without generating a PPT')

    args = parser.parse_args()

    # Ensure the output filename has a .pptx extension
    output_filename = args.output
    if not output_filename.lower().endswith('.pptx'):
        output_filename += '.pptx'

    # Get the PPT name without its extension
    ppt_name = os.path.splitext(os.path.basename(output_filename))[0]

    # Build the PPT project directory path
    ppt_project_dir = os.path.join(args.base_dir, ppt_name)
    origin_image_dir = os.path.join(ppt_project_dir, "origin_image")

    if args.init:
        os.makedirs(origin_image_dir, exist_ok=True)
        print(f"✓ PPT project directory ready: {ppt_project_dir}")
        print(f"✓ Slide image directory ready: {origin_image_dir}")
        sys.exit(0)

    if not os.path.exists(ppt_project_dir):
        print(f"Error: PPT project directory does not exist: {ppt_project_dir}")
        print("Add --init to initialize the directories")
        sys.exit(1)

    if not os.path.exists(origin_image_dir):
        print(f"Error: origin_image directory does not exist: {origin_image_dir}")
        print("Add --init to initialize the directories")
        sys.exit(1)

    # Set the output path
    output_path = os.path.join(ppt_project_dir, output_filename)

    # Get slide images
    print(f"Scanning PPT project directory: {ppt_project_dir}")
    image_files = get_slide_images(ppt_project_dir)

    if not image_files:
        print("Error: no image files found")
        print("Supported formats: .png, .jpg, .jpeg, .gif, .bmp")
        print(f"\nPlace slide images in: {origin_image_dir}/")
        sys.exit(1)

    print(f"Found {len(image_files)} slide images\n")
    speaker_notes = load_speaker_notes(ppt_project_dir)
    if speaker_notes:
        print(f"Found notes for {len(speaker_notes)} slides: {os.path.join(ppt_project_dir, 'speech.md')}\n")

    # Create the presentation
    print(f"Creating PPT (aspect ratio: {args.aspect_ratio})...")
    print("-" * 50)

    success = create_presentation(image_files, output_path, args.aspect_ratio, speaker_notes)

    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
